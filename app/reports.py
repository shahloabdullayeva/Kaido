import csv
import io
from datetime import date, datetime, timedelta

from .config import ROOT
from .db import rows
from .queries import service_status, trucks_with_service
from .work_types import GROUP_LABELS, GROUP_OF, KIND_LABELS

FONTS = ROOT / "app" / "fonts"


REPORTS = {
    "maintenance": "Maintenance and repairs",
    "fuel": "Fuel purchases",
    "breakdowns": "Breakdowns",
    "faults": "Fault codes",
    "fleet": "Fleet service status",
}


def _scope(filters, alias="t"):
    clauses, params = [], []
    if filters.get("truck_id"):
        clauses.append(f"{alias}.id = %s")
        params.append(filters["truck_id"])
    if filters.get("fleet") == "own":
        clauses.append(f"not {alias}.is_outside")
    elif filters.get("fleet") == "outside":
        clauses.append(f"{alias}.is_outside")
    return clauses, params


def _where(company_col, filters, date_expr=None, driver_col=None, alias="t"):
    clauses = [f"{company_col} = %s"]
    params = [filters["company_id"]]
    if date_expr and filters.get("start"):
        clauses.append(f"{date_expr} >= %s")
        params.append(filters["start"])
    if date_expr and filters.get("end"):
        clauses.append(f"{date_expr} < %s")
        params.append(filters["end"] + timedelta(days=1))
    if driver_col and filters.get("driver_id"):
        clauses.append(f"{driver_col} = %s")
        params.append(filters["driver_id"])
    extra, extra_params = _scope(filters, alias)
    return " and ".join(clauses + extra), params + extra_params


def maintenance(filters):
    where, params = _where("m.company_id", filters, "coalesce(m.performed_on, m.scheduled_for)", "m.driver_id")
    data = rows(
        f"""select m.id, coalesce(m.performed_on, m.scheduled_for) as day, t.unit_number,
              case when t.is_outside then coalesce(t.outside_carrier, 'outside') else '' end as carrier,
              d.name as driver, m.kind, m.status, m.vendor, m.invoice_no, m.odometer, m.cost, m.description
            from maintenance_orders m join trucks t on t.id = m.truck_id
            left join drivers d on d.id = m.driver_id
            where {where} order by day desc nulls last, m.id desc""",
        params,
    )
    for row in data:
        row["system"] = GROUP_LABELS.get(GROUP_OF.get(row["kind"]), "")
        row["kind"] = KIND_LABELS.get(row["kind"], row["kind"])
        row["status"] = row["status"].replace("_", " ")
    columns = [("id", "#", "int", 11), ("day", "Date", "date", 18), ("unit_number", "Unit", "text", 16),
               ("carrier", "Outside carrier", "text", 22), ("driver", "Driver", "text", 26),
               ("system", "System", "text", 22), ("kind", "Work", "text", 26), ("status", "Status", "text", 14), ("vendor", "Shop", "text", 30),
               ("invoice_no", "Invoice", "text", 14), ("odometer", "Odometer", "int", 17),
               ("cost", "Cost", "money", 14), ("description", "What was done", "long", 60)]
    totals = {"cost": sum(float(r["cost"] or 0) for r in data)}
    return columns, data, totals


def fuel(filters):
    where, params = _where("f.company_id", filters, "f.purchased_at", "f.driver_id")
    data = rows(
        f"""select f.purchased_at as day, t.unit_number, d.name as driver, f.fuel_type, f.gallons,
              f.price_per_gallon, f.total, f.odometer, f.location, f.state, f.card_last4, f.source
            from fuel_transactions f join trucks t on t.id = f.truck_id
            left join drivers d on d.id = f.driver_id
            where {where} order by f.purchased_at desc""",
        params,
    )
    columns = [("day", "Date", "datetime", 20), ("unit_number", "Unit", "text", 14), ("driver", "Driver", "text", 26),
               ("fuel_type", "Type", "text", 10), ("gallons", "Gallons", "num", 12),
               ("price_per_gallon", "Per gallon", "money3", 12), ("total", "Total", "money", 14),
               ("odometer", "Odometer", "int", 14), ("location", "Where", "text", 36), ("state", "State", "text", 8),
               ("card_last4", "Card", "text", 8), ("source", "Source", "text", 10)]
    totals = {"gallons": sum(float(r["gallons"] or 0) for r in data), "total": sum(float(r["total"] or 0) for r in data)}
    return columns, data, totals


def breakdowns(filters):
    where, params = _where("b.company_id", filters, "b.occurred_at", "b.driver_id")
    data = rows(
        f"""select b.id, b.occurred_at as day, t.unit_number, d.name as driver, b.severity, b.status, b.location,
              case when b.resolved_at is not null
                   then round(extract(epoch from (b.resolved_at - b.occurred_at)) / 3600.0, 1) end as down_hours,
              b.towing_cost, b.repair_cost, b.description, b.cause, b.resolution
            from breakdowns b join trucks t on t.id = b.truck_id
            left join drivers d on d.id = b.driver_id
            where {where} order by b.occurred_at desc""",
        params,
    )
    for row in data:
        row["status"] = row["status"].replace("_", " ")
    columns = [("id", "#", "int", 8), ("day", "Reported", "datetime", 20), ("unit_number", "Unit", "text", 14),
               ("driver", "Driver", "text", 24), ("severity", "Severity", "text", 10), ("status", "Status", "text", 12),
               ("location", "Where", "text", 30), ("down_hours", "Hours down", "num", 12),
               ("towing_cost", "Towing", "money", 12), ("repair_cost", "Repair", "money", 12),
               ("description", "What happened", "long", 50), ("resolution", "What fixed it", "long", 40)]
    totals = {"towing_cost": sum(float(r["towing_cost"] or 0) for r in data),
              "repair_cost": sum(float(r["repair_cost"] or 0) for r in data)}
    return columns, data, totals


def faults(filters):
    where, params = _where("f.company_id", filters, "f.first_seen_at")
    data = rows(
        f"""select f.first_seen_at as day, t.unit_number,
              coalesce(f.dtc_code, 'SPN ' || f.spn || ' / FMI ' || f.fmi) as code,
              f.description, f.lamp, f.severity, f.occurrence_count, f.last_seen_at, f.cleared_at
            from fault_events f join trucks t on t.id = f.truck_id
            where {where} order by f.first_seen_at desc""",
        params,
    )
    columns = [("day", "First seen", "datetime", 20), ("unit_number", "Unit", "text", 12), ("code", "Code", "text", 20),
               ("description", "Description", "long", 50), ("lamp", "Lamp", "text", 12),
               ("severity", "Severity", "text", 10), ("occurrence_count", "Times", "int", 8),
               ("last_seen_at", "Last seen", "datetime", 20), ("cleared_at", "Cleared", "datetime", 20)]
    return columns, data, {}


def fleet(filters):
    data = []
    for truck in trucks_with_service(filters["company_id"]):
        if truck["status"] == "sold":
            continue
        if filters.get("truck_id") and truck["id"] != filters["truck_id"]:
            continue
        if filters.get("fleet") == "own" and truck["is_outside"]:
            continue
        if filters.get("fleet") == "outside" and not truck["is_outside"]:
            continue
        status = {item["label"]: item for item in service_status(truck)}
        data.append({
            "unit_number": truck["unit_number"],
            "vehicle": " ".join(str(p) for p in [truck["year"], truck["make"], truck["model"]] if p),
            "engine": truck.get("engine"),
            "driver": truck.get("driver_name"),
            "status": truck["status"].replace("_", " "),
            "odometer": truck["odometer"],
            "last_oil": truck.get("last_oil_odometer"),
            "oil": status["Oil service"]["detail"],
            "dot": status["DOT annual"]["detail"],
            "other": "; ".join(f"{item['label']}: {item['detail']}" for item in status.values()
                               if not item.get("core") and item["tone"] in ("bad", "warn")),
            "faults": truck.get("active_faults"),
            "breakdowns": truck.get("open_breakdowns"),
        })
    columns = [("unit_number", "Unit", "text", 12), ("vehicle", "Vehicle", "text", 28), ("engine", "Engine", "text", 22),
               ("driver", "Driver", "text", 26), ("status", "Status", "text", 12), ("odometer", "Odometer", "int", 14),
               ("last_oil", "Last oil at", "int", 14), ("oil", "Oil service", "text", 22), ("dot", "DOT annual", "text", 20),
               ("other", "Other service due", "text", 40),
               ("faults", "Active faults", "int", 10), ("breakdowns", "Open breakdowns", "int", 10)]
    return columns, data, {}


BUILDERS = {"maintenance": maintenance, "fuel": fuel, "breakdowns": breakdowns, "faults": faults, "fleet": fleet}


def build(kind, filters):
    return BUILDERS[kind](filters)


def cell_text(value, kind):
    if value is None or value == "":
        return ""
    if kind == "money":
        return "${:,.2f}".format(float(value))
    if kind == "money3":
        return "${:,.3f}".format(float(value))
    if kind == "int":
        return "{:,}".format(int(value))
    if kind == "num":
        return "{:,.1f}".format(float(value))
    if kind == "date":
        return value.strftime("%d %b %Y") if hasattr(value, "strftime") else str(value)
    if kind == "datetime":
        return value.strftime("%d %b %Y %H:%M") if hasattr(value, "strftime") else str(value)
    return str(value)


def to_csv(columns, data, totals):
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow([label for _, label, _, _ in columns])
    for row in data:
        writer.writerow([_plain(row.get(key)) for key, _, _, _ in columns])
    if totals:
        writer.writerow([("Total" if index == 0 else (f"{totals[key]:.2f}" if key in totals else ""))
                         for index, (key, _, _, _) in enumerate(columns)])
    return out.getvalue().encode("utf-8-sig")


def _plain(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    return value


def _excel_value(value):
    if isinstance(value, datetime) and value.tzinfo:
        return value.replace(tzinfo=None)
    if value is not None and not isinstance(value, (int, float, str, date, datetime)):
        return float(value)
    return value


def to_xlsx(title, subtitle, columns, data, totals):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = title[:31]
    sheet["A1"] = title
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = subtitle
    sheet["A2"].font = Font(italic=True, color="555555")
    header_row = 4
    thin = Side(style="thin", color="999999")
    for index, (_, label, _, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=header_row, column=index, value=label)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
        cell.border = Border(bottom=thin)
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width, len(label) + 2), 70)
    formats = {"money": '"$"#,##0.00', "money3": '"$"#,##0.000', "int": "#,##0", "num": "#,##0.0",
               "date": "dd mmm yyyy", "datetime": "dd mmm yyyy hh:mm"}
    for offset, row in enumerate(data, start=1):
        for index, (key, _, kind, _) in enumerate(columns, start=1):
            cell = sheet.cell(row=header_row + offset, column=index, value=_excel_value(row.get(key)))
            if kind in formats:
                cell.number_format = formats[kind]
            if kind == "long":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    last = header_row + len(data)
    if totals and data:
        total_row = last + 1
        sheet.cell(row=total_row, column=1, value="Total").font = Font(bold=True)
        for index, (key, _, kind, _) in enumerate(columns, start=1):
            if key in totals:
                letter = get_column_letter(index)
                cell = sheet.cell(row=total_row, column=index, value=f"=SUM({letter}{header_row + 1}:{letter}{last})")
                cell.font = Font(bold=True)
                cell.number_format = formats.get(kind, "#,##0.00")
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    if data:
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{last}"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _pdf():
    from fpdf import FPDF
    pdf = FPDF(orientation="L", unit="mm", format="Letter")
    pdf.add_font("dejavu", "", str(FONTS / "DejaVuSansCondensed.ttf"))
    pdf.add_font("dejavu", "B", str(FONTS / "DejaVuSansCondensed-Bold.ttf"))
    pdf.set_auto_page_break(True, margin=12)
    pdf.set_margins(10, 10, 10)
    return pdf


def to_pdf(title, subtitle, columns, data, totals):
    pdf = _pdf()
    pdf.set_title(title)
    pdf.add_page()
    pdf.set_font("dejavu", "B", 15)
    pdf.cell(0, 8, title, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("dejavu", "", 9)
    pdf.set_text_color(90, 90, 90)
    pdf.cell(0, 5, subtitle, new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0, 0, 0)
    pdf.ln(3)
    shown = [c for c in columns if c[2] != "long"] + [c for c in columns if c[2] == "long"][:1]
    widths = [c[3] for c in shown]
    scale = pdf.epw / sum(widths)
    widths = [w * scale for w in widths]
    right = {"money", "money3", "int", "num"}
    from fpdf.fonts import FontFace
    bold = FontFace(emphasis="BOLD")
    pdf.set_font("dejavu", "", 7.5)
    with pdf.table(col_widths=widths, text_align=["RIGHT" if c[2] in right else "LEFT" for c in shown],
                   line_height=4.2, first_row_as_headings=True, repeat_headings=1,
                   headings_style=FontFace(emphasis="BOLD", fill_color=(221, 221, 221))) as table:
        header = table.row()
        for column in shown:
            header.cell(column[1])
        for row in data:
            line = table.row()
            for key, _, kind, _ in shown:
                text = cell_text(row.get(key), kind)
                line.cell(text[:400] if kind == "long" else text)
        if totals and data:
            line = table.row()
            for index, (key, _, kind, _) in enumerate(shown):
                line.cell("Total" if index == 0 else (cell_text(totals[key], kind) if key in totals else ""), style=bold)
    if not data:
        pdf.set_font("dejavu", "", 10)
        pdf.cell(0, 8, "Nothing in this period.")
    return bytes(pdf.output())


def record_pdf(title, subtitle, fields, sections, tables=()):
    pdf = _pdf()
    pdf.set_title(title)
    pdf.add_page(orientation="P")
    pdf.set_font("dejavu", "B", 16)
    pdf.cell(0, 9, title, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("dejavu", "", 9)
    pdf.set_text_color(90, 90, 90)
    pdf.cell(0, 5, subtitle, new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0, 0, 0)
    pdf.ln(4)
    from fpdf.fonts import FontFace
    bold = FontFace(emphasis="BOLD")
    pdf.set_font("dejavu", "", 9)
    with pdf.table(col_widths=(45, 145), first_row_as_headings=False, line_height=5.5) as table:
        for label, value in fields:
            line = table.row()
            line.cell(label, style=bold)
            line.cell(value or "—")
    for heading, widths, header, body in tables:
        if not body:
            continue
        pdf.ln(4)
        pdf.set_font("dejavu", "B", 11)
        pdf.cell(0, 7, heading, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("dejavu", "", 9)
        with pdf.table(col_widths=widths, line_height=5.5, text_align=("LEFT", "LEFT", "RIGHT", "RIGHT", "RIGHT")) as grid:
            top = grid.row()
            for text in header:
                top.cell(text, style=bold)
            for values in body:
                line = grid.row()
                for value in values:
                    line.cell(value or "")
    for heading, body in sections:
        if not body:
            continue
        pdf.ln(4)
        pdf.set_font("dejavu", "B", 11)
        pdf.cell(0, 7, heading, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("dejavu", "", 9.5)
        pdf.multi_cell(0, 5, body)
    pdf.ln(8)
    pdf.set_font("dejavu", "", 8)
    pdf.set_text_color(120, 120, 120)
    pdf.cell(0, 5, f"Printed from Kaido {datetime.now():%d %b %Y %H:%M}")
    return bytes(pdf.output())
