import csv
import io
import re
from datetime import date, datetime, timedelta


def norm(header):
    return re.sub(r"[^a-z0-9]+", " ", str(header or "").lower()).strip()


def read(filename, data):
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        return _read_xlsx(data)
    if name.endswith(".xls"):
        raise ValueError("Old .xls files cannot be read. Open it and save it as .xlsx or .csv.")
    return _read_csv(data)


def _read_csv(data):
    text = None
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    table = [row for row in csv.reader(io.StringIO(text), dialect) if any(cell.strip() for cell in row)]
    return _rows_from_table(table)


def _read_xlsx(data):
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = book.worksheets[0]
    table = []
    for row in sheet.iter_rows(values_only=True):
        cells = ["" if cell is None else cell for cell in row]
        if any(str(cell).strip() for cell in cells):
            table.append(cells)
    book.close()
    return _rows_from_table(table)


def _rows_from_table(table):
    if not table:
        return [], []
    header_index = 0
    for index, row in enumerate(table[:15]):
        filled = [cell for cell in row if str(cell).strip()]
        if len(filled) >= 3 and sum(1 for cell in filled if re.search(r"[A-Za-z]", str(cell))) >= len(filled) * 0.6:
            header_index = index
            break
    headers = [norm(cell) for cell in table[header_index]]
    out = []
    for row in table[header_index + 1:]:
        record = {}
        for position, header in enumerate(headers):
            if header and position < len(row):
                record.setdefault(header, row[position])
        out.append(record)
    return headers, out


def find(headers, *candidates):
    for wanted in candidates:
        if wanted in headers:
            return wanted
    for wanted in candidates:
        for header in headers:
            if header.startswith(wanted) or wanted in header.split():
                return header
    return None


def value(record, header):
    if not header:
        return None
    raw = record.get(header)
    if raw is None:
        return None
    if isinstance(raw, str):
        raw = raw.strip()
        return raw or None
    return raw


def number(raw):
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        return float(raw)
    text = str(raw).strip().replace("$", "").replace(",", "")
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    try:
        result = float(text)
    except ValueError:
        return None
    return -result if negative else result


def when(raw, time_raw=None):
    if raw is None:
        return None
    if isinstance(raw, datetime):
        moment = raw
    elif isinstance(raw, date):
        moment = datetime(raw.year, raw.month, raw.day)
    elif isinstance(raw, (int, float)) and 20000 < raw < 80000:
        moment = datetime(1899, 12, 30) + timedelta(days=float(raw))
    else:
        text = str(raw).strip()
        moment = None
        for fmt in ("%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M", "%m/%d/%Y %I:%M %p", "%m/%d/%Y",
                    "%m/%d/%y %H:%M", "%m/%d/%y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S",
                    "%Y-%m-%d %H:%M", "%Y-%m-%d", "%d.%m.%Y", "%d %b %Y"):
            try:
                moment = datetime.strptime(text, fmt)
                break
            except ValueError:
                continue
        if moment is None:
            return None
    if time_raw is not None and moment.hour == 0 and moment.minute == 0:
        if isinstance(time_raw, datetime):
            moment = moment.replace(hour=time_raw.hour, minute=time_raw.minute)
        elif hasattr(time_raw, "hour"):
            moment = moment.replace(hour=time_raw.hour, minute=time_raw.minute)
        else:
            for fmt in ("%H:%M:%S", "%H:%M", "%I:%M %p", "%I:%M:%S %p"):
                try:
                    parsed = datetime.strptime(str(time_raw).strip(), fmt)
                    moment = moment.replace(hour=parsed.hour, minute=parsed.minute)
                    break
                except ValueError:
                    continue
    return moment


def unit_key(raw):
    if raw is None:
        return ""
    if isinstance(raw, float) and raw.is_integer():
        raw = int(raw)
    text = str(raw).strip().lower().replace("unit", "").replace("#", "").strip()
    text = re.sub(r"\s+", "", text)
    stripped = text.lstrip("0")
    return stripped or text


def unit_lead(raw):
    if raw is None:
        return ""
    if isinstance(raw, float) and raw.is_integer():
        raw = int(raw)
    found = re.match(r"\s*(?:unit)?\s*#?\s*(\d+[a-z]?)(?![a-z0-9])", str(raw), re.I)
    return unit_key(found.group(1)) if found else ""
